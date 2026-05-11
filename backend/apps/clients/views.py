from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from rest_framework.pagination import PageNumberPagination
from django.db.models import Q
from openpyxl import load_workbook
import csv
from io import BytesIO, StringIO

from .models import Client
from .serializers import ClientSerializer, ClientUploadSerializer
from apps.companies.models import Company
from apps.users.permissions import HasPermission


class ClientPagination(PageNumberPagination):
    page_size = 50
    page_size_query_param = 'page_size'
    max_page_size = 500


class ClientViewSet(viewsets.ModelViewSet):
    """ViewSet for Client CRUD operations"""
    
    queryset = Client.objects.all()
    serializer_class = ClientSerializer
    pagination_class = ClientPagination
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = 'clients'
    
    def get_queryset(self):
        """Filter clients based on query parameters"""
        queryset = Client.objects.all()
        
        # Search filter
        search = self.request.query_params.get('search', None)
        if search:
            queryset = queryset.filter(
                Q(client_code__icontains=search) |
                Q(full_name__icontains=search) |
                Q(pan_or_citizenship__icontains=search) |
                Q(boid__icontains=search)
            )
        
        # Status filter
        status_filter = self.request.query_params.get('status', None)
        if status_filter:
            queryset = queryset.filter(status=status_filter)
        
        # Holder type filter
        holder_type = self.request.query_params.get('holder_type', None)
        if holder_type:
            queryset = queryset.filter(holder_type=holder_type)

        # Company filter (by id)
        company = self.request.query_params.get('company', None)
        if company:
            queryset = queryset.filter(company_id=company)
        
        return queryset
    
    @action(detail=False, methods=['post'])
    def upload(self, request):
        """Bulk upload clients from Excel/CSV"""
        serializer = ClientUploadSerializer(data=request.data)
        
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        
        file = serializer.validated_data['file']
        
        try:
            # Read file
            data = []
            if file.name.endswith('.csv'):
                reader = csv.DictReader(StringIO(file.read().decode('utf-8')))
                data = list(reader)
            else:
                wb = load_workbook(BytesIO(file.read()))
                sheet = wb.active
                headers = [cell.value for cell in sheet[1]]
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    data.append(dict(zip(headers, row)))
            
            # Expected columns (BOID is required)
            expected_columns = [
                'client_code', 'full_name', 'boid', 'holder_type',
                'company', 'pan_or_citizenship', 'bank_name', 'bank_account_no'
            ]
            
            if not data:
                return Response({'error': 'No data found in file'}, status=status.HTTP_400_BAD_REQUEST)
            
            # Validate required columns (client_code, full_name, boid)
            missing_columns = [col for col in expected_columns[:3] if col not in data[0]]
            if missing_columns:
                return Response(
                    {'error': f'Missing required columns: {", ".join(missing_columns)}'},
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            # Process data
            created = 0
            updated = 0
            errors = []
            
            for index, row in enumerate(data):
                try:
                    client_code = str(row['client_code']).strip().upper()
                    full_name = str(row['full_name']).strip()
                    boid = str(row.get('boid', '')).strip().upper()
                    
                    if not client_code or not full_name or not boid:
                        errors.append(f"Row {index + 2}: Client code, name and BOID are required")
                        continue
                    
                    client_data = {
                        'client_code': client_code,
                        'full_name': full_name,
                        'boid': boid,
                        'holder_type': str(row.get('holder_type', '')).strip() or None,
                        'pan_or_citizenship': str(row.get('pan_or_citizenship', '')).strip() or None,
                        'bank_name': str(row.get('bank_name', '')).strip() or None,
                        'bank_account_no': str(row.get('bank_account_no', '')).strip() or None,
                    }
                    # handle company if provided (code or name)
                    company_val = str(row.get('company', '')).strip()
                    if company_val:
                        comp = Company.objects.filter(
                            Q(company_code__iexact=company_val) |
                            Q(company_name__iexact=company_val)
                        ).first()
                        if comp:
                            client_data['company'] = comp
                    
                    # Update or create
                    client, created_flag = Client.objects.update_or_create(
                        client_code=client_code,
                        defaults=client_data
                    )
                    
                    if created_flag:
                        created += 1
                    else:
                        updated += 1
                        
                except Exception as e:
                    errors.append(f"Row {index + 2}: {str(e)}")
            
            return Response({
                'message': 'Upload completed',
                'created': created,
                'updated': updated,
                'errors': errors
            }, status=status.HTTP_200_OK)
            
        except Exception as e:
            return Response(
                {'error': f'Error processing file: {str(e)}'},
                status=status.HTTP_400_BAD_REQUEST
            )

    @action(detail=False, methods=['get'])
    def lookup_boid(self, request):
        """Lookup client by BOID"""
        boid = request.query_params.get('boid')
        if not boid:
            return Response({'error': 'BOID is required as query parameter'}, status=status.HTTP_400_BAD_REQUEST)
        try:
            client = Client.objects.get(boid__iexact=boid.strip())
            serializer = ClientSerializer(client)
            return Response(serializer.data)
        except Client.DoesNotExist:
            return Response({'message': 'Not found'}, status=status.HTTP_404_NOT_FOUND)
        from django.http import HttpResponse
        import xlsxwriter
        
        output = BytesIO()
        workbook = xlsxwriter.Workbook(output)
        worksheet = workbook.add_worksheet('Clients')
        
        # Headers (include BOID)
        headers = [
            'client_code', 'full_name', 'boid', 'holder_type',
            'company', 'pan_or_citizenship', 'bank_name', 'bank_account_no'
        ]
        
        # Write headers
        header_format = workbook.add_format({'bold': True, 'bg_color': '#D9E1F2'})
        for col, header in enumerate(headers):
            worksheet.write(0, col, header, header_format)
        
        # Sample data (include BOID sample)
        sample_data = [
            ['CL001', 'Ram Kumar Shrestha', 'BOID-CL001', 'Public', 'COMP001', '12345678', 'NIC Asia Bank', '1234567890123'],
            ['CL002', 'ABC Investment Pvt. Ltd', 'BOID-CL002', 'Institution', 'COMP002', '987654321', 'Standard Chartered', '9876543210987'],
        ]
        
        for row_idx, row_data in enumerate(sample_data, start=1):
            for col_idx, cell_data in enumerate(row_data):
                worksheet.write(row_idx, col_idx, cell_data)
        
        workbook.close()
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=client_upload_template.xlsx'
        
        return response
