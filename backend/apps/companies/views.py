from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django.db.models import Q
from openpyxl import load_workbook
import csv
from io import BytesIO, StringIO

from .models import Company
from .serializers import CompanySerializer, CompanyUploadSerializer
from apps.users.permissions import HasPermission


class CompanyViewSet(viewsets.ModelViewSet):
    """ViewSet for Company CRUD operations"""
    
    queryset = Company.objects.all()
    serializer_class = CompanySerializer
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = 'companies'
    
    def get_queryset(self):
        """Filter companies based on query parameters"""
        queryset = Company.objects.all()
        
        # Search filter
        search = self.request.query_params.get('search', None)
        if search:
            queryset = queryset.filter(
                Q(company_code__icontains=search) |
                Q(company_name__icontains=search) |
                Q(pan_no__icontains=search)
            )
        
        # Status filter
        status_filter = self.request.query_params.get('status', None)
        if status_filter:
            queryset = queryset.filter(status=status_filter)
        
        # Sector filter (support legacy and new param names)
        sector = self.request.query_params.get('sector') or self.request.query_params.get('sector_type')
        if sector:
            queryset = queryset.filter(sector_type=sector)
        
        # Tax status filter
        tax_status = self.request.query_params.get('tax_status') or self.request.query_params.get('interest_tax_status')
        if tax_status:
            queryset = queryset.filter(interest_tax_status=tax_status)
        
        return queryset
    
    @action(detail=False, methods=['post'])
    def upload(self, request):
        """Bulk upload companies from Excel/CSV"""
        serializer = CompanyUploadSerializer(data=request.data)
        
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        
        file = serializer.validated_data['file']
        
        try:
            # Read file based on extension
            data = []
            if file.name.endswith('.csv'):
                reader = csv.DictReader(StringIO(file.read().decode('utf-8')))
                data = list(reader)
            else:
                wb = load_workbook(BytesIO(file.read()))
                sheet = wb.active
                headers = [cell.value for cell in sheet[1]]
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    data.append(dict(zip(headers, row)))
            
            # Expected columns
            expected_columns = [
                'company_code', 'company_name', 'sector_type', 
                'interest_tax_status', 'pan_no', 'bank_name', 'bank_account_no'
            ]
            
            if not data:
                return Response({'error': 'No data found in file'}, status=status.HTTP_400_BAD_REQUEST)
            
            # Validate columns
            missing_columns = [col for col in expected_columns[:2] if col not in data[0]]
            if missing_columns:
                return Response(
                    {'error': f'Missing required columns: {", ".join(missing_columns)}'},
                    status=status.HTTP_400_BAD_REQUEST
                )
            
            # Process data
            created = 0
            updated = 0
            errors = []
            
            for index, row in enumerate(data):
                try:
                    company_code = str(row['company_code']).strip().upper()
                    company_name = str(row['company_name']).strip()
                    
                    if not company_code or not company_name:
                        errors.append(f"Row {index + 2}: Company code and name are required")
                        continue
                    
                    company_data = {
                        'company_code': company_code,
                        'company_name': company_name,
                        'sector_type': str(row.get('sector_type', '')).strip() or None,
                        'interest_tax_status': str(row.get('interest_tax_status', '')).strip() or None,
                        'pan_no': str(row.get('pan_no', '')).strip() or None,
                        'bank_name': str(row.get('bank_name', '')).strip() or None,
                        'bank_account_no': str(row.get('bank_account_no', '')).strip() or None,
                    }
                    
                    # Check if company exists
                    company, created_flag = Company.objects.update_or_create(
                        company_code=company_code,
                        defaults=company_data
                    )
                    
                    if created_flag:
                        created += 1
                    else:
                        updated += 1
                        
                except Exception as e:
                    errors.append(f"Row {index + 2}: {str(e)}")
            
            return Response({
                'message': 'Upload completed',
                'created': created,
                'updated': updated,
                'errors': errors
            }, status=status.HTTP_200_OK)
            
        except Exception as e:
            return Response(
                {'error': f'Error processing file: {str(e)}'},
                status=status.HTTP_400_BAD_REQUEST
            )
    
    @action(detail=False, methods=['get'])
    def export_template(self, request):
        """Download Excel template for company upload"""
        from django.http import HttpResponse
        import xlsxwriter
        from io import BytesIO
        
        output = BytesIO()
        workbook = xlsxwriter.Workbook(output)
        worksheet = workbook.add_worksheet('Companies')
        
        # Headers
        headers = [
            'company_code', 'company_name', 'sector_type',
            'interest_tax_status', 'pan_no', 'bank_name', 'bank_account_no'
        ]
        
        # Write headers
        header_format = workbook.add_format({'bold': True, 'bg_color': '#D9E1F2'})
        for col, header in enumerate(headers):
            worksheet.write(0, col, header, header_format)
        
        # Sample data
        sample_data = [
            ['COMP001', 'Sample Company Ltd', 'Public', 'Taxable', '123456789', 'Nepal Bank Ltd', '1234567890'],
            ['COMP002', 'Another Company Pvt. Ltd', 'Private', 'Exempted', '987654321', 'Himalayan Bank', '0987654321'],
        ]
        
        for row_idx, row_data in enumerate(sample_data, start=1):
            for col_idx, cell_data in enumerate(row_data):
                worksheet.write(row_idx, col_idx, cell_data)
        
        workbook.close()
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=company_upload_template.xlsx'
        
        return response
