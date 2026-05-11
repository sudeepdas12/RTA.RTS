import csv
from datetime import datetime
from io import BytesIO, StringIO

from openpyxl import load_workbook

from apps.clients.models import Client
from apps.companies.models import Company


def load_upload_rows(file_obj):
    """Read CSV/XLSX into list of dictionaries."""
    if file_obj.name.endswith('.csv'):
        reader = csv.DictReader(StringIO(file_obj.read().decode('utf-8')))
        return list(reader)

    workbook = load_workbook(BytesIO(file_obj.read()))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    data = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        data.append(dict(zip(headers, row)))
    return data


def missing_columns(data, required_cols):
    if not data:
        return required_cols
    return [column for column in required_cols if column not in data[0]]


def to_float(value, default=0.0):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def parse_date(value, formats=None):
    if not value:
        return None

    accepted_formats = formats or ('%Y-%m-%d', '%m/%d/%Y')
    raw = str(value).strip()
    for fmt in accepted_formats:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


def normalize_payment_status(raw_status, valid_statuses):
    status = str(raw_status or '').strip().title() or 'Pending'
    return status if status in valid_statuses else None


def resolve_client(row, row_index, errors):
    """Resolve client by code or BOID and append row-level errors when needed."""
    row_no = row_index + 2

    if row.get('client_code'):
        client_code = str(row['client_code']).strip().upper()
        try:
            return Client.objects.get(client_code=client_code)
        except Client.DoesNotExist:
            errors.append(f"Row {row_no}: Client code {client_code} not found")
            return None

    if row.get('boid'):
        boid = str(row['boid']).strip().upper()
        try:
            return Client.objects.get(boid__iexact=boid)
        except Client.DoesNotExist:
            errors.append(f"Row {row_no}: BOID {boid} not found")
            return None

    errors.append(f"Row {row_no}: client_code or boid is required")
    return None


def resolve_company(company_code, row_index, errors):
    row_no = row_index + 2
    try:
        return Company.objects.get(company_code=company_code)
    except Company.DoesNotExist:
        errors.append(f"Row {row_no}: Company {company_code} not found")
        return None
