from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from openpyxl import load_workbook
import csv
from io import BytesIO, StringIO
from decimal import Decimal
from datetime import datetime

from .models import BankStatement, BankTransaction, Reconciliation
from .serializers import (
    BankStatementSerializer, BankTransactionSerializer,
    ReconciliationSerializer, BankStatementUploadSerializer
)
from apps.payables.models import InterestPayable, DividendPayable
from apps.users.permissions import HasPermission


class BankStatementViewSet(viewsets.ModelViewSet):
    queryset = BankStatement.objects.all()
    serializer_class = BankStatementSerializer
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = 'reconciliation'
    
    @action(detail=False, methods=['post'])
    def upload(self, request):
        """Upload bank statement"""
        serializer = BankStatementUploadSerializer(data=request.data)
        
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
        
        file = serializer.validated_data['file']
        
        try:
            # Create bank statement record
            user = getattr(request, 'user_obj', None) or request.user
            bank_stmt = BankStatement.objects.create(
                bank_name=serializer.validated_data['bank_name'],
                account_no=serializer.validated_data['account_no'],
                statement_from=serializer.validated_data['statement_from'],
                statement_to=serializer.validated_data['statement_to'],
                uploaded_by=user,
                file_name=file.name
            )
            
            # Read file
            data = []
            if file.name.endswith('.csv'):
                reader = csv.DictReader(StringIO(file.read().decode('utf-8')))
                data = list(reader)
            else:
                wb = load_workbook(BytesIO(file.read()))
                sheet = wb.active
                headers = [cell.value for cell in sheet[1]]
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    data.append(dict(zip(headers, row)))
            
            # Process transactions
            total_debit = 0
            total_credit = 0
            created = 0
            
            for index, row in enumerate(data):
                try:
                    txn_date_str = str(row['txn_date'])
                    try:
                        txn_date = datetime.strptime(txn_date_str, '%Y-%m-%d').date()
                    except ValueError:
                        try:
                            txn_date = datetime.strptime(txn_date_str, '%m/%d/%Y').date()
                        except ValueError:
                            continue  # skip invalid date
                    
                    debit = float(row.get('debit', 0) or 0)
                    credit = float(row.get('credit', 0) or 0)
                    
                    BankTransaction.objects.create(
                        bank_stmt=bank_stmt,
                        txn_date=txn_date,
                        reference_no=str(row.get('reference_no', '')).strip() or None,
                        debit=debit,
                        credit=credit,
                        balance=float(row.get('balance', 0) or 0),
                        description=str(row.get('description', '')).strip() or None
                    )
                    
                    total_debit += debit
                    total_credit += credit
                    created += 1
                except Exception as e:
                    pass
            
            # Update totals
            bank_stmt.total_debit = total_debit
            bank_stmt.total_credit = total_credit
            bank_stmt.save()
            
            return Response({
                'message': 'Bank statement uploaded successfully',
                'bank_stmt_id': bank_stmt.bank_stmt_id,
                'transactions_created': created
            }, status=status.HTTP_201_CREATED)
            
        except Exception as e:
            return Response(
                {'error': f'Error processing file: {str(e)}'},
                status=status.HTTP_400_BAD_REQUEST
            )
    
    @action(detail=False, methods=['get'])
    def export_template(self, request):
        """Download Excel template for bank statement upload"""
        from django.http import HttpResponse
        import xlsxwriter
        
        output = BytesIO()
        workbook = xlsxwriter.Workbook(output)
        worksheet = workbook.add_worksheet('Bank Transactions')
        
        headers = [
            'txn_date', 'reference_no', 'description',
            'debit', 'credit', 'balance'
        ]
        
        header_format = workbook.add_format({'bold': True, 'bg_color': '#D9E1F2'})
        for col, header in enumerate(headers):
            worksheet.write(0, col, header, header_format)
        
        sample_data = [
            ['2026-02-01', 'REF001', 'Interest payment for COMP001', 0, 42500, 1000000],
            ['2026-02-02', 'REF002', 'Dividend payment', 0, 106250, 1106250],
            ['2026-02-03', 'REF003', 'Service charges', 500, 0, 1105750],
        ]
        
        for row_idx, row_data in enumerate(sample_data, start=1):
            for col_idx, cell_data in enumerate(row_data):
                worksheet.write(row_idx, col_idx, cell_data)
        
        workbook.close()
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=bank_statement_template.xlsx'
        
        return response


class BankTransactionViewSet(viewsets.ReadOnlyModelViewSet):
    queryset = BankTransaction.objects.all()
    serializer_class = BankTransactionSerializer
    permission_classes = [IsAuthenticated]
    
    def get_queryset(self):
        queryset = BankTransaction.objects.all()
        
        # Filter by statement
        bank_stmt_id = self.request.query_params.get('bank_stmt', None)
        if bank_stmt_id:
            queryset = queryset.filter(bank_stmt_id=bank_stmt_id)
        
        return queryset


class ReconciliationViewSet(viewsets.ModelViewSet):
    queryset = Reconciliation.objects.all()
    serializer_class = ReconciliationSerializer
    permission_classes = [IsAuthenticated, HasPermission]
    required_permission = 'reconciliation'
    
    def perform_create(self, serializer):
        user = getattr(self.request, 'user_obj', None) or self.request.user
        serializer.save(reconciled_by=user)
    
    @action(detail=False, methods=['post'])
    def auto_match(self, request):
        """Auto-match bank transactions with payables"""
        bank_stmt_id = request.data.get('bank_stmt_id')
        
        if not bank_stmt_id:
            return Response(
                {'error': 'bank_stmt_id is required'},
                status=status.HTTP_400_BAD_REQUEST
            )
        
        try:
            bank_stmt = BankStatement.objects.get(bank_stmt_id=bank_stmt_id)
        except BankStatement.DoesNotExist:
            return Response(
                {'error': 'Bank statement not found'},
                status=status.HTTP_404_NOT_FOUND
            )
        
        transactions = bank_stmt.transactions.all()
        matched = 0

        # Track already reconciled source records to avoid duplicate matching of the same payable.
        interest_already_linked_ids = set(
            Reconciliation.objects.filter(source_type='Interest').values_list('source_id', flat=True)
        )
        dividend_already_linked_ids = set(
            Reconciliation.objects.filter(source_type='Dividend').values_list('source_id', flat=True)
        )
        
        for txn in transactions:
            # Skip if already reconciled
            if txn.reconciliations.exists():
                continue
            
            amount = txn.credit if txn.credit > 0 else txn.debit
            
            # Try to match with interest payables
            interest_match = InterestPayable.objects.filter(
                net_payable=amount,
                payment_status='Pending'
            ).exclude(interest_id__in=interest_already_linked_ids).first()
            
            if interest_match:
                user = getattr(request, 'user_obj', None) or request.user
                Reconciliation.objects.create(
                    bank_txn=txn,
                    source_type='Interest',
                    source_id=interest_match.interest_id,
                    matched_amount=amount,
                    recon_status='Matched',
                    reconciled_by=user
                )
                interest_already_linked_ids.add(interest_match.interest_id)
                matched += 1
                continue
            
            # Try to match with dividend payables
            dividend_match = DividendPayable.objects.filter(
                net_payable=amount,
                payment_status='Pending'
            ).exclude(dividend_id__in=dividend_already_linked_ids).first()
            
            if dividend_match:
                user = getattr(request, 'user_obj', None) or request.user
                Reconciliation.objects.create(
                    bank_txn=txn,
                    source_type='Dividend',
                    source_id=dividend_match.dividend_id,
                    matched_amount=amount,
                    recon_status='Matched',
                    reconciled_by=user
                )
                dividend_already_linked_ids.add(dividend_match.dividend_id)
                matched += 1
        
        return Response({
            'message': 'Auto-matching completed',
            'matched_count': matched,
            'total_transactions': transactions.count()
        })
