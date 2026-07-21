"""
Import PRIME debenture data from the Excel file.
Reads the ORIGINAL sheet (raw data with BOID, name, kitta, bank details), 
calculates interest for Shrawan 082 to Poush End 082 (182 days @ 8.75%),
stores in debenture_reconciliation table, making it visible in the 
Debenture Interest Reconciliation page.
"""
import os, sys

os.environ['DJANGO_SETTINGS_MODULE'] = 'config.settings'
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import django
django.setup()

from django.db import transaction
from django.db.models import Sum, Count
from datetime import date
import uuid


def import_from_excel(excel_path, company_code='PRIME', company_name='PRIME'):
    """Import debenture data from PRIME Excel file"""
    try:
        import openpyxl
    except ImportError:
        print("ERROR: openpyxl not installed. Run: pip install openpyxl")
        sys.exit(1)

    from apps.payables.models import DebentureReconciliation
    
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    print(f"Sheets available: {wb.sheetnames}")
    
    # --- Read ORIGINAL sheet (raw bank data with KITTA) ---
    if 'ORIGINAL' not in wb.sheetnames:
        print("ERROR: ORIGINAL sheet not found!")
        return
    
    ws = wb['ORIGINAL']
    print(f"\n--- ORIGINAL Sheet: {ws.max_row} rows, {ws.max_column} cols ---")
    
    # Interest calculation parameters (from PUBLIC sheet header)
    interest_rate = 8.75
    tax_rate = 6.00
    period_from = date(2025, 7, 17)  # 1 Shrawan 082
    period_to = date(2026, 1, 14)    # Poush End 082
    period_days = 182
    
    # Also try to read period info from PUBLIC sheet
    if 'PUBLIC' in wb.sheetnames:
        ws_pub = wb['PUBLIC']
        for r_idx in range(1, min(6, ws_pub.max_row + 1)):
            row_vals = []
            for c in range(1, min(8, ws_pub.max_column + 1)):
                row_vals.append(ws_pub.cell(row=r_idx, column=c).value)
            print(f"  PUBLIC Row {r_idx}: {row_vals}")
    
    # Build records from ORIGINAL sheet
    records = []
    header_row = None
    
    # Find header row
    for r_idx in range(1, min(11, ws.max_row + 1)):
        first_val = ws.cell(row=r_idx, column=1).value
        if first_val and str(first_val).strip().upper() in ('S.NO', 'S.N', 'SN', 'S.N.'):
            header_row = r_idx
            break
    
    if header_row is None:
        print("ERROR: Could not find header row in ORIGINAL sheet")
        return
    
    print(f"Header row found at row {header_row}")
    
    # Read headers
    headers = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=header_row, column=c).value
        if val:
            headers[str(val).strip().upper()] = c
    
    print(f"Headers found: {list(headers.keys())}")
    
    # Map columns
    def hdr(*names):
        for n in names:
            for k, v in headers.items():
                if n.upper() in k:
                    return v
        return None
    
    col_sn = hdr('S.NO', 'S.N', 'SN')
    col_boid = hdr('BOID', 'B/O')
    col_name = hdr('NAME', 'APPLICANT', 'SHAREHOLDER')
    col_father = hdr('FATHERS NAME', 'FATHER')
    col_grandfather = hdr('GRANDFATHERS NAME', 'GRANDFATHER', 'SPOUSE')
    col_citizenship = hdr('CITIZENSHIP', 'CITIZEN')
    col_district = hdr('DISTRICT')
    col_bank_code = hdr('BANK CODE', 'BANK_CODE')
    col_bank_name = hdr('BANK NAME', 'BANK')
    col_bank_acct = hdr('BANK ACCOUNT', 'ACCOUNT', 'ACCOUNT NUMBER', 'BANK AC')
    col_kitta = hdr('KITTA', 'ALLOTED', 'QUANTITY', 'UNITS', 'SHARES')
    
    print(f"\nColumn mapping:")
    print(f"  S.N: {col_sn}, BOID: {col_boid}, NAME: {col_name}")
    print(f"  FATHER: {col_father}, GRANDFATHER: {col_grandfather}")
    print(f"  CITIZENSHIP: {col_citizenship}, DISTRICT: {col_district}")
    print(f"  BANK CODE: {col_bank_code}, BANK NAME: {col_bank_name}")
    print(f"  BANK ACCT: {col_bank_acct}, KITTA: {col_kitta}")
    
    if not col_sn:
        print("ERROR: Could not find S.N column!")
        return
    
    # Process each row
    for r_idx in range(header_row + 1, ws.max_row + 1):
        sn_val = ws.cell(row=r_idx, column=col_sn).value
        
        if sn_val is None:
            continue
        
        try:
            sn = int(float(str(sn_val).strip()))
        except (ValueError, TypeError):
            continue
        
        def get_str(col_idx):
            if col_idx is None:
                return ""
            v = ws.cell(row=r_idx, column=col_idx).value
            if v is None:
                return ""
            return str(v).strip()
        
        kitta = 0
        if col_kitta:
            kitta_raw = ws.cell(row=r_idx, column=col_kitta).value
            if kitta_raw is not None:
                try:
                    kitta = float(str(kitta_raw).replace(',', ''))
                except (ValueError, TypeError):
                    kitta = 0
        
        amount = kitta * 1000  # Each kitta = NPR 1000
        
        # Calculate interest
        # Gross Interest = Amount * Rate% * Days/365
        gross_interest = amount * (interest_rate / 100) * period_days / 365
        
        # TDS @ tax_rate%
        tax_amount = gross_interest * (tax_rate / 100)
        
        # Net interest
        net_interest = gross_interest - tax_amount
        
        # Roundup (Net Payable = net interest payable rounded to 2 decimal places)
        roundup = round(net_interest, 2)
        
        records.append({
            'sn': sn,
            'boid': get_str(col_boid),
            'applicant_name': get_str(col_name),
            'father_mother_name': get_str(col_father),
            'grandfather_spouse_name': get_str(col_grandfather),
            'citizenship_number': get_str(col_citizenship),
            'issued_from': get_str(col_district),
            'alloted_quantity': kitta,
            'amount': amount,
            'annual_interest': round(amount * interest_rate / 100, 4),
            'daily_interest': round(amount * interest_rate / 100 / 365, 6),
            'period_interest': round(gross_interest, 4),
            'tax': round(tax_amount, 4),
            'net_interest_payable': round(net_interest, 4),
            'roundup': roundup,
            'bank_code': get_str(col_bank_code),
            'bank_name': get_str(col_bank_name),
            'account_number': get_str(col_bank_acct),
            'lot': 'PRIME LOT 1',
            'status': 'SUCCESS',
            'remarks': '',
        })
    
    print(f"\nTotal records extracted from ORIGINAL sheet: {len(records)}")
    
    # Show first 3 records as sample
    print("\nSample records:")
    for r in records[:3]:
        print(f"  SN={r['sn']}, NAME={r['applicant_name'][:30]}, BOID={r['boid']}, KITTA={r['alloted_quantity']}, AMOUNT={r['amount']}")
        print(f"    GROSS_INT={r['period_interest']}, TAX={r['tax']}, NET={r['net_interest_payable']}, ROUNDUP={r['roundup']}")
        print(f"    BANK={r['bank_name']}({r['bank_code']}), ACCT={r['account_number']}")
    
    # Save to DB
    print(f"\n--- Saving to database ---")
    upload_batch = str(uuid.uuid4())[:8]
    user = None
    try:
        from django.contrib.auth import get_user_model
        user = get_user_model().objects.filter(is_superuser=True).first()
    except Exception:
        pass
    
    batch_size = 500
    created = 0
    
    # Clear existing data for this company first
    existing = DebentureReconciliation.objects.filter(company_code=company_code).count()
    if existing > 0:
        print(f"Found {existing} existing records for {company_code}. Deleting...")
        DebentureReconciliation.objects.filter(company_code=company_code).delete()
    
    with transaction.atomic():
        for i in range(0, len(records), batch_size):
            batch = records[i:i+batch_size]
            objs = []
            for r in batch:
                objs.append(DebentureReconciliation(
                    company_code=company_code,
                    company_name=company_name,
                    report_title=f"{interest_rate}% {company_name} Debenture Interest Reconciliation",
                    report_subtitle=f"From {period_from} to {period_to} ({period_days} days)",
                    sector_type='Public',
                    tax_status='Taxable',
                    sn=r['sn'],
                    boid=r['boid'],
                    applicant_name=r['applicant_name'],
                    father_mother_name=r['father_mother_name'],
                    grandfather_spouse_name=r['grandfather_spouse_name'],
                    citizenship_number=r['citizenship_number'],
                    issued_from=r['issued_from'],
                    alloted_quantity=r['alloted_quantity'],
                    amount=r['amount'],
                    annual_interest=r['annual_interest'],
                    daily_interest=r['daily_interest'],
                    period_interest=r['period_interest'],
                    tax=r['tax'],
                    net_interest_payable=r['net_interest_payable'],
                    roundup=r['roundup'],
                    bank_code=r['bank_code'],
                    bank_name=r['bank_name'],
                    account_number=r['account_number'],
                    lot=r['lot'],
                    status=r['status'],
                    remarks=r['remarks'],
                    period_from=period_from,
                    period_to=period_to,
                    period_days=period_days,
                    interest_rate=interest_rate,
                    tax_rate=tax_rate,
                    upload_batch=upload_batch,
                    created_by=user,
                ))
            DebentureReconciliation.objects.bulk_create(objs, ignore_conflicts=False)
            created += len(objs)
            print(f"  Created {created}/{len(records)}...")
    
    print(f"\n{'='*60}")
    print(f"SUCCESS: {created} records imported for {company_name} ({company_code})")
    print(f"{'='*60}")
    
    # Verify
    total = DebentureReconciliation.objects.filter(company_code=company_code).count()
    stats = DebentureReconciliation.objects.filter(company_code=company_code).aggregate(
        total_amount=Sum('amount'),
        total_gross=Sum('period_interest'),
        total_tax=Sum('tax'),
        total_net=Sum('roundup'),
    )
    print(f"Total in DB for {company_code}: {total} records")
    for k, v in stats.items():
        val = v or 0
        print(f"  {k}: NPR {float(val):,.2f}")


if __name__ == '__main__':
    excel_path = r"C:\Users\Administrator\OneDrive\Desktop\PRIME 8.75 POUSH END 2082_CALCULATION - FINAL.xlsx"
    
    if not os.path.exists(excel_path):
        print(f"ERROR: File not found at {excel_path}")
        alt_paths = [
            "PRIME 8.75 POUSH END 2082_CALCULATION - FINAL.xlsx",
            r"C:\Users\Administrator\Downloads\PRIME 8.75 POUSH END 2082_CALCULATION - FINAL.xlsx",
        ]
        for p in alt_paths:
            if os.path.exists(p):
                excel_path = p
                break
        else:
            print("Please provide the correct path to the Excel file")
            sys.exit(1)
    
    import_from_excel(excel_path, company_code='PRIME', company_name='Prime Commercial Bank Ltd.')